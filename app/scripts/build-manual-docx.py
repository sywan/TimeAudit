"""Build the printable Time Audit manuals from their Markdown sources.

Run with Python 3 and python-docx. Convert the resulting DOCX files to PDF
with LibreOffice; the PDF files shipped on GitHub Pages are generated from
these DOCX sources.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PUBLIC = Path(__file__).resolve().parent.parent / "public"
BLUE = RGBColor(54, 105, 137)
MUTED = RGBColor(95, 119, 135)


def set_font(style, western, east_asian, size, color=None):
    style.font.name = western
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, value in (("ascii", western), ("hAnsi", western), ("eastAsia", east_asian)):
        rfonts.set(qn(f"w:{key}"), value)


def add_inline(paragraph, raw):
    for part in re.split(r"(\*\*[^*]+\*\*|\[[^]]+\]\([^)]+\))", raw):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("[") and "](" in part:
            label, url = part[1:-1].split("](", 1)
            paragraph.add_run(f"{label}: {url}")
        else:
            paragraph.add_run(part)


def build(language):
    name = "USER-MANUAL.md" if language == "zh-TW" else "USER-MANUAL.en.md"
    source = (PUBLIC / name).read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = Cm(2.15)
    section.bottom_margin = Cm(1.9)
    section.left_margin = section.right_margin = Cm(2.35)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.05)

    east = "Arial Unicode MS" if language == "zh-TW" else "Arial"
    western = east if language == "zh-TW" else "Arial"
    normal = doc.styles["Normal"]
    set_font(normal, western, east, 10.5 if language == "zh-TW" else 10.25)
    normal.paragraph_format.line_spacing = 1.36
    normal.paragraph_format.space_after = Pt(6)

    for key, size, before, after in (("Title", 23, 0, 13), ("Heading 1", 14, 17, 7)):
        style = doc.styles[key]
        set_font(style, western, east, size, RGBColor(0, 0, 0))
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    set_font(doc.styles["List Bullet"], western, east, 10.5)
    set_font(doc.styles["List Number"], western, east, 10.5)
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("TIME AUDIT  /  USER GUIDE")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = BLUE

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("V1.20  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

    lines = source.splitlines()
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_inline(p, line[2:])
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:])
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        elif re.match(r"\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", line))
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
            if line.startswith("©") or line.startswith("Interactive version:") or line.startswith("線上互動版："):
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = MUTED

    doc.core_properties.title = "Time Audit 使用手冊" if language == "zh-TW" else "Time Audit User Guide"
    doc.core_properties.author = "Shu-Yen Wan"
    output = PUBLIC / ("USER-MANUAL.zh-TW.docx" if language == "zh-TW" else "USER-MANUAL.en.docx")
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build("zh-TW")
    build("en")
